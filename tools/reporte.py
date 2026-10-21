"""Reportes diarios de vencimientos por gerente (Excel para OneDrive).

Uso:  python3 reporte.py <ultimo.json o '-'> <aplicativo index.html> <destinatarios.json> <carpeta_salida>
- ultimo.json: cierre que envía la hoja de Google a GitHub ('-' = sin datos, genera los archivos vacíos).
- aplicativo: página del aplicativo (sucursales, unidad, zona y SKU con inventario por tienda).
- destinatarios.json: [{"correo", "nombre", "archivo", "alcance": {"unidad": ...} | {"zona": ...} | {"todas": true}}]
Imprime JSON con el resumen por destinatario (lotes, unidades críticas) para armar los correos.
"""
import sys, json, datetime as dt, os, re, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, FormulaRule

src, app, dest, out = sys.argv[1:5]
os.makedirs(out, exist_ok=True)
page = open(app, encoding='utf-8').read()
D = json.JSONDecoder().raw_decode(page[page.index('const D = ') + 10:])[0]
datos = json.load(open(src, encoding='utf-8')) if src != '-' else {'cierre': (dt.date.today() - dt.timedelta(days=1)).isoformat(), 'filas': [], 'avance': {}}
cierre = dt.date.fromisoformat(datos['cierre']); hoy = cierre + dt.timedelta(days=1)
iy, iw, _ = cierre.isocalendar(); semana = f'{iy}-S{iw:02d}'
corto = lambda s: re.sub(r'^[^|]+\| ', '', s)

# Último reporte por sucursal + código (todas sus filas con la misma marca de tiempo)
grupos = collections.defaultdict(list)
for f in datos['filas']: grupos[(f[3], f[5])].append(f)
lotes = []
for (suc, cod), fs in grupos.items():
  ult = max(f[13] for f in fs)
  for f in fs:
    if f[13] != ult: continue
    v = dt.date.fromisoformat(f[10][:10]); dias = (v - hoy).days
    lotes.append({'suc': suc, 'marca': f[4], 'cod': cod, 'prod': f[6], 'inv': f[7], 'sit': f[8], 'u': f[9], 'v': v, 'dias': dias,
                  'col': f[11], 'reg': dt.datetime.fromisoformat(f[13]), 'nuevo': f[13][:10] == cierre.isoformat()})
estado = lambda d: 'Vencido' if d < 0 else 'Crítico ≤30' if d <= 30 else 'Alto 31–60' if d <= 60 else 'Medio 61+'

FONT = 'Arial'; H = PatternFill('solid', fgColor='D4161E')
def libro(dst):
  al = dst['alcance']
  sucs = [i for i, s in enumerate(D['s']) if al.get('todas') or s[0] == al.get('unidad') or s[1] == al.get('zona')]
  names = {D['s'][i][2]: i for i in sucs}
  L = sorted([l for l in lotes if l['suc'] in names], key=lambda l: (D['s'][names[l['suc']]][0], D['s'][names[l['suc']]][1], corto(l['suc']), l['dias'], l['prod']))
  wb = Workbook()
  # ---------- Detalle
  ws = wb.active; ws.title = 'Detalle'
  ws['A1'] = f"Vencimientos en tienda · {dst['nombre']} · cierre al {cierre:%d/%m/%Y}"; ws['A1'].font = Font(name=FONT, bold=True, size=13)
  ws['A2'] = f"Último reporte de cada producto por sucursal (últimos 35 días). Días para vencer contados al {hoy:%d/%m/%Y}. Usa las flechas de los títulos para filtrar."
  ws['A2'].font = Font(name=FONT, size=9, italic=True, color='5A4F50')
  cols = ['Unidad', 'Zona', 'Sucursal', 'Marca', 'Código', 'Producto', 'Inv. sistema (u)', 'Por vencer (u)', 'Vencimiento', 'Días para vencer', 'Estado', 'Situación', 'Colaborador', 'Fecha de reporte', 'Nuevo ayer']
  ws.append([]); ws.append(cols)
  for l in L:
    s = D['s'][names[l['suc']]]
    ws.append([s[0], corto(s[1]), corto(s[2]), l['marca'], l['cod'], l['prod'], l['inv'], l['u'], l['v'], l['dias'], estado(l['dias']), l['sit'], l['col'], l['reg'], 'Sí' if l['nuevo'] else 'No'])
  last = max(ws.max_row, 5)
  if ws.max_row == 4: ws.append(['Sin lotes reportados'] + [''] * (len(cols) - 1)); last = 5
  for r in ws.iter_rows(min_row=5, max_row=last):
    for c in r: c.font = Font(name=FONT, size=10)
    r[4].number_format = '@'; r[8].number_format = 'dd/mm/yyyy'; r[13].number_format = 'dd/mm/yyyy hh:mm'
  t = Table(displayName='Detalle', ref=f'A4:O{last}'); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium3', showRowStripes=True)
  ws.add_table(t)
  ws.conditional_formatting.add(f'K5:K{last}', FormulaRule(formula=['OR($K5="Vencido",$K5="Crítico ≤30")'], fill=PatternFill('solid', fgColor='F8E0E0'), font=Font(color='B42828', bold=True)))
  ws.conditional_formatting.add(f'K5:K{last}', CellIsRule(operator='equal', formula=['"Alto 31–60"'], fill=PatternFill('solid', fgColor='FBE6DC'), font=Font(color='A3431C')))
  ws.conditional_formatting.add(f'K5:K{last}', CellIsRule(operator='equal', formula=['"Medio 61+"'], fill=PatternFill('solid', fgColor='FDF1D4'), font=Font(color='8A5A00')))
  for col, w in zip('ABCDEFGHIJKLMNO', [17, 10, 24, 18, 14, 40, 11, 11, 12, 10, 13, 12, 18, 16, 9]): ws.column_dimensions[col].width = w
  ws.freeze_panes = 'D5'
  # ---------- Resumen por sucursal
  rs = wb.create_sheet('Resumen por sucursal')
  rs['A1'] = f"Avance de la semana {iw} y unidades por vencer · cierre al {cierre:%d/%m/%Y}"; rs['A1'].font = Font(name=FONT, bold=True, size=13)
  rc = ['Unidad', 'Zona', 'Sucursal', 'SKU con inventario', 'Revisados semana', '% avance', 'SKU con vencimiento', 'Vencido (u)', '≤30 días (u)', '31–60 días (u)', '61+ días (u)', 'Último registro', 'Registró']
  rs.append([]); rs.append(rc)
  tot = {'lotes': 0, 'crit': 0}
  for i in sorted(sucs, key=lambda i: (D['s'][i][0], D['s'][i][1], D['s'][i][2])):
    s = D['s'][i]; a = datos['avance'].get(f'{semana}|{s[2]}', {})
    ll = [l for l in L if l['suc'] == s[2]]; b = [sum(l['u'] for l in ll if c(l['dias'])) for c in (lambda d: d < 0, lambda d: 0 <= d <= 30, lambda d: 31 <= d <= 60, lambda d: d > 60)]
    rev = a.get('n', 0) + a.get('e', 0); r = rs.max_row + 1
    rs.append([s[0], corto(s[1]), corto(s[2]), len(D['i'][i]) // 2, rev, f'=IFERROR(E{r}/D{r},0)', len({l['cod'] for l in ll}), *b,
               dt.datetime.fromisoformat(a['act']) if a.get('act') else None, (a.get('quien') or '').split(' · ')[0]])
    tot['lotes'] += len(ll); tot['crit'] += b[0] + b[1]
  lr = rs.max_row
  for r in rs.iter_rows(min_row=4, max_row=lr):
    for c in r: c.font = Font(name=FONT, size=10)
    r[5].number_format = '0%'; r[11].number_format = 'dd/mm/yyyy hh:mm'
  rt = Table(displayName='Resumen', ref=f'A3:M{lr}'); rt.tableStyleInfo = TableStyleInfo(name='TableStyleMedium3', showRowStripes=True); rs.add_table(rt)
  rs.cell(lr + 1, 3, 'Total').font = Font(name=FONT, bold=True)
  for c, L_ in zip(range(4, 12), 'DEFGHIJK'):
    rs.cell(lr + 1, c, f'=IFERROR(E{lr+1}/D{lr+1},0)' if L_ == 'F' else f'=SUBTOTAL(109,{L_}4:{L_}{lr})').font = Font(name=FONT, bold=True)
  rs.cell(lr + 1, 6).number_format = '0%'
  for col, w in zip('ABCDEFGHIJKLM', [17, 10, 26, 12, 12, 10, 12, 11, 11, 11, 11, 16, 18]): rs.column_dimensions[col].width = w
  rs.freeze_panes = 'D4'
  wb.save(os.path.join(out, dst['archivo']))
  return {'correo': dst['correo'], 'archivo': dst['archivo'], 'sucursales': len(sucs), 'lotes': tot['lotes'], 'u_criticas': tot['crit'],
          'nuevos_ayer': sum(1 for l in L if l['nuevo'])}

res = [libro(d) for d in json.load(open(dest, encoding='utf-8'))]
print(json.dumps({'cierre': cierre.isoformat(), 'semana': semana, 'destinatarios': res}, ensure_ascii=False))
